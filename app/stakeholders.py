"""4. 사업관리 › 이해관계자 — 행·열을 늘리고 줄일 수 있는 명단 + 사진·비고 + 엑셀 내려받기 (v0.7)

· 기본 열(BASE_COLS)은 앱이 쓰는 열이라 지울 수 없고, [열 추가]로 만든 열만 [열 삭제]로 지운다.
· 사진은 pmc/4_사업관리/이해관계자_사진/ 에 저장하고 photo_path 열에 상대 경로를 적는다 (문서 폴더 밖이라 AI 검색에 섞이지 않는다).
· 엑셀은 openpyxl 로 만들고, 사진은 Pillow 가 있으면 칸 안에 작게 넣는다.
"""
import io
import re
from datetime import date
from pathlib import Path

from app import storage
from app.db import ROOT, connect

TABLE = "stakeholders"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BASE_COLS = ["stakeholder_id", "category", "org", "dept", "person_name", "position", "phone", "email", "valid_from", "valid_to", "note", "photo_path"]
HIDDEN = {"contact"}                                       # DROP COLUMN 이 안 되는 옛 SQLite 에서만 남아 있는 열
EDIT_COLS = ["stakeholder_id", "category", "org", "dept", "person_name", "position", "phone", "email", "valid_from", "valid_to", "note"]
LABELS = {"stakeholder_id": "관계자ID", "category": "구분", "org": "기관", "dept": "부서", "person_name": "담당자", "position": "직위",
          "phone": "휴대전화", "email": "이메일", "valid_from": "시작일", "valid_to": "종료일", "note": "비고", "photo_path": "사진"}
PHOTO_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}


def label(c: str) -> str:
    return LABELS.get(c, c)


def extra_columns() -> list[str]:
    """[열 추가]로 만든 사용자 열 (DB 열 순서대로)"""
    return [c for c in storage.columns(TABLE) if c not in BASE_COLS and c not in HIDDEN]


def rows() -> list[dict]:
    with connect(readonly=True) as con:
        cur = con.execute(f"SELECT * FROM {TABLE} ORDER BY stakeholder_id")
        cols = [d[0] for d in cur.description]
        out = [dict(zip(cols, r)) for r in cur.fetchall()]
    for r in out:
        for k, v in r.items():
            r[k] = "" if v is None else v
    return out


def next_id(existing: list[str] | None = None) -> str:
    ids = existing if existing is not None else [r["stakeholder_id"] for r in rows()]
    nums = [int(m.group(1)) for i in ids for m in [re.search(r"(\d+)\s*$", str(i))] if m]
    return f"SH-{(max(nums) + 1 if nums else 1):02d}"


def fill_ids(df):
    """편집표에서 비워 둔 관계자ID 채우기"""
    df = df.copy()
    ids = [str(x) for x in df["stakeholder_id"].fillna("").astype(str) if str(x).strip()]
    for i in df.index:
        if not str(df.at[i, "stakeholder_id"] or "").strip():
            new = next_id(ids)
            df.at[i, "stakeholder_id"] = new
            ids.append(new)
    return df


def save(df) -> tuple[bool, str]:
    """편집표 저장 — 편집표에 없는 열(사진 경로, 감춘 열)은 DB 값을 그대로 유지한다"""
    df = fill_ids(df)
    keep = {r["stakeholder_id"]: r for r in rows()}
    for c in [c for c in storage.columns(TABLE) if c not in df.columns]:
        df[c] = [keep.get(str(i), {}).get(c, "") for i in df["stakeholder_id"]]
    ok, msg = storage.save_table(TABLE, df)
    if ok:
        _prune_photos()
    return ok, msg


# ───────────────────────── 사진 ─────────────────────────
def photo_dir() -> Path:
    p = storage.folder("4_사업관리") / "이해관계자_사진"
    p.mkdir(parents=True, exist_ok=True)
    return p


def photo_file(r: dict) -> Path | None:
    rel = str(r.get("photo_path") or "").strip()
    if not rel:
        return None
    p = ROOT / rel
    return p if p.exists() else None


def save_photo(sid: str, filename: str, data: bytes) -> tuple[bool, str]:
    storage.check_edit()
    ext = Path(filename).suffix.lower()
    if ext not in PHOTO_EXTS:
        return False, "jpg·png·gif·webp 그림 파일만 올릴 수 있습니다."
    if len(data) > 15 * 1024 * 1024:
        return False, "15MB 이하 파일만 올릴 수 있습니다."
    r = next((x for x in rows() if x["stakeholder_id"] == sid), None)
    if r is None:
        return False, "먼저 [저장하기]로 행을 저장한 뒤 사진을 올려 주세요."
    delete_photo(sid, update=False)
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", str(r.get("person_name") or r.get("org") or "")).strip()[:30]
    p = photo_dir() / f"{sid}_{name}{ext}".replace(" ", "_")
    p.write_bytes(data)
    return storage.update_cell(TABLE, "stakeholder_id", sid, "photo_path", storage.rel(p).replace("\\", "/"))


def delete_photo(sid: str, update: bool = True) -> None:
    storage.check_edit()
    r = next((x for x in rows() if x["stakeholder_id"] == sid), None)
    f = photo_file(r) if r else None
    if f:
        try:
            f.unlink()
        except OSError:
            pass
    if update:
        storage.update_cell(TABLE, "stakeholder_id", sid, "photo_path", "")


def _prune_photos() -> None:
    """지워진 행의 사진 파일 정리"""
    used = {str(f.resolve()) for r in rows() for f in [photo_file(r)] if f}
    for f in photo_dir().glob("*"):
        if f.is_file() and str(f.resolve()) not in used:
            try:
                f.unlink()
            except OSError:
                pass


# ───────────────────────── 엑셀 ─────────────────────────
def build_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cols = [c for c in EDIT_COLS if c != "note"] + extra_columns() + ["note", "photo_path"]
    data = rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "이해관계자"
    ws["A1"] = f"이해관계자 명단 — KOICA 필리핀 루존 부카스센터 PMC  (출력 {date.today():%Y-%m-%d}, {len(data)}명)"
    ws["A1"].font = Font(name="맑은 고딕", size=12, bold=True, color="1F4E9C")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    head_fill, thin = PatternFill("solid", fgColor="EEF2F9"), Side(style="thin", color="D5DCE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=2, column=j, value=label(c))
        cell.font, cell.fill, cell.border = Font(name="맑은 고딕", bold=True, size=10), head_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {"stakeholder_id": 10, "category": 12, "org": 18, "dept": 24, "person_name": 12, "position": 12, "phone": 15, "email": 26,
              "valid_from": 11, "valid_to": 11, "note": 36, "photo_path": 14}
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 16)

    try:
        from PIL import Image as PILImage
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        PILImage = XLImage = None

    photo_col = cols.index("photo_path") + 1
    for i, r in enumerate(data, 3):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=i, column=j, value="" if c == "photo_path" else r.get(c, ""))
            cell.font, cell.border = Font(name="맑은 고딕", size=10), border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        f = photo_file(r)
        if f and PILImage and XLImage:
            try:
                im = PILImage.open(f).convert("RGB")
                im.thumbnail((110, 110))
                buf = io.BytesIO()
                im.save(buf, format="PNG")
                buf.seek(0)
                img = XLImage(buf)
                img.width, img.height = im.size
                ws.add_image(img, f"{get_column_letter(photo_col)}{i}")
                ws.row_dimensions[i].height = max(ws.row_dimensions[i].height or 15, im.size[1] * 0.78 + 6)
                ws.column_dimensions[get_column_letter(photo_col)].width = max(ws.column_dimensions[get_column_letter(photo_col)].width, im.size[0] / 7 + 2)
            except Exception:
                ws.cell(row=i, column=photo_col, value=f.name)
        elif f:
            ws.cell(row=i, column=photo_col, value=f.name)
    ws.freeze_panes = "A3"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
