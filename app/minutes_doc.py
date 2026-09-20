"""회의록 1건 → Word(.docx) / PDF / Markdown(검색용) 파일 만들기 (v0.5)

· Word  : python-docx  (맑은 고딕, 표 형식 회의록 양식)
· PDF   : reportlab    (윈도우 맑은 고딕 → 없으면 Noto CJK → 없으면 내장 한글 CID 폰트 순으로 사용)
· 화면의 [회의록 추가]에서 저장 전에도 다운로드할 수 있도록 모두 bytes 를 반환한다.
"""
import html
import io
import re
from datetime import datetime
from pathlib import Path

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
PROJECT = "KOICA 필리핀 루존지역 부카스센터 긴급 외래 진료 및 진단검사 역량강화 사업 — PMC"

# (표시 이름, 회의록 열)
META_FIELDS = [("회의명", "title"), ("일시", "_when"), ("장소", "location"), ("대분류 (Activity)", "activity"),
               ("중분류 (구분)", "meeting_type"), ("소분류 (이해관계기관)", "stakeholder_org"), ("참석자", "participants"), ("작성자", "author")]
SECTIONS = [("1. 회의 주요 안건", "agenda"), ("2. 회의 결과 요약", "summary"), ("3. 후속조치사항", "follow_up")]


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def when_text(m: dict) -> str:
    return (_s(m.get("meeting_date")) + " " + _s(m.get("meeting_time"))).strip()


def file_stem(m: dict) -> str:
    t = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", _s(m.get("title")))[:50].strip() or "회의록"
    return f"회의록_{_s(m.get('meeting_date')) or '날짜미정'}_{t}"


def _meta(m: dict) -> list[tuple[str, str]]:
    m = dict(m)
    m["_when"] = when_text(m)
    return [(k, _s(m.get(c))) for k, c in META_FIELDS]


def _action_lines(actions: list[dict] | None) -> list[str]:
    out = []
    for a in actions or []:
        line = _s(a.get("description"))
        extra = " · ".join(x for x in (f"담당 {_s(a.get('owner'))}" if _s(a.get("owner")) else "",
                                       f"기한 {_s(a.get('due_date'))}" if _s(a.get("due_date")) else "",
                                       _s(a.get("status"))) if x)
        out.append(line + (f"  ({extra})" if extra else ""))
    return out


# ───────────────────────── Markdown (AI 사업비서 검색용) ─────────────────────────
def to_markdown(m: dict, actions: list[dict] | None = None) -> str:
    head = ["---", f"doc_id: {_s(m.get('doc_id')) or 'MTG_' + _s(m.get('meeting_id'))}", f"title: \"{_s(m.get('title')).replace(chr(34), chr(39))}\"",
            "doc_type: 회의록", f"doc_date: {_s(m.get('meeting_date'))}", f"source_org: \"{_s(m.get('stakeholder_org')).replace(chr(34), chr(39))}\"", "---", ""]
    body = [f"# 회의록 — {_s(m.get('title'))}", ""]
    body += [f"- {k}: {v}" for k, v in _meta(m) if v]
    body.append(f"- 회의ID: {_s(m.get('meeting_id'))}")
    for title, col in SECTIONS:
        body += ["", f"## {title}", "", _s(m.get(col)) or "(없음)"]
    lines = _action_lines(actions)
    if lines:
        body += ["", "### 등록된 후속조치", ""] + [f"- {x}" for x in lines]
    return "\n".join(head + body) + "\n"


# ───────────────────────── Word ─────────────────────────
def build_docx(m: dict, actions: list[dict] | None = None) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "맑은 고딕", Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    def _font(run, size=None, bold=None, color=None):
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def _shade(cell, hex_color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    p = doc.add_paragraph()
    _font(p.add_run(PROJECT), 9, color="5B6675")
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(h.add_run("회 의 록"), 20, True, "1F4E9C")
    h.paragraph_format.space_after = Pt(10)

    t = doc.add_table(rows=0, cols=4)
    t.style, t.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
    meta = _meta(m)

    def _pad(cell):
        pf = cell.paragraphs[0].paragraph_format
        pf.space_before, pf.space_after = Pt(3), Pt(3)

    def _kv(cell_k, cell_v, k, v, bold_v=False):
        _shade(cell_k, "EEF2F9")
        _pad(cell_k); _pad(cell_v)
        _font(cell_k.paragraphs[0].add_run(k), 10, True)
        _font(cell_v.paragraphs[0].add_run(v), 10.5, bold_v)

    first = t.add_row().cells                              # 회의명은 한 줄 전체
    merged = first[1].merge(first[3])
    _kv(first[0], merged, meta[0][0], meta[0][1], True)
    rest = meta[1:]
    for i in range(0, len(rest), 2):
        row = t.add_row().cells
        for j, (k, v) in enumerate(rest[i:i + 2]):
            _kv(row[j * 2], row[j * 2 + 1], k, v)
    for r in t.rows:
        r.cells[0].width = r.cells[2].width = Cm(3.3)
        r.cells[1].width = r.cells[3].width = Cm(5.2)

    for title, col in SECTIONS:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before, hp.paragraph_format.space_after = Pt(14), Pt(4)
        _font(hp.add_run(title), 13, True, "1F4E9C")
        text = _s(m.get(col)) or "(없음)"
        for line in text.splitlines() or [""]:
            bp = doc.add_paragraph()
            bp.paragraph_format.space_after = Pt(2)
            _font(bp.add_run(line if line.strip() else " "), 11)
    lines = _action_lines(actions)
    if lines:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(8)
        _font(hp.add_run("등록된 후속조치 (현황)"), 11, True)
        at = doc.add_table(rows=1, cols=4)
        at.style = "Table Grid"
        for c, k in zip(at.rows[0].cells, ["조치 내용", "담당", "기한", "상태"]):
            _shade(c, "EEF2F9"); _pad(c); _font(c.paragraphs[0].add_run(k), 10, True)
        for a in actions:
            cells = at.add_row().cells
            for c, k in zip(cells, ["description", "owner", "due_date", "status"]):
                _pad(c); _font(c.paragraphs[0].add_run(_s(a.get(k))), 10)
        for r in at.rows:
            for c, w in zip(r.cells, (9.3, 3.2, 2.6, 2.0)):
                c.width = Cm(w)
    f = doc.add_paragraph()
    f.paragraph_format.space_before = Pt(18)
    _font(f.add_run(f"회의ID {_s(m.get('meeting_id')) or '(저장 전)'} · 출력 {datetime.now():%Y-%m-%d %H:%M} · KOICA AI 성과관리"), 8.5, color="8A97A8")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ───────────────────────── PDF ─────────────────────────
_FONT_CANDIDATES = [  # (이름, 경로, TTC 안의 순번)
    ("MalgunGothic", r"C:\Windows\Fonts\malgun.ttf", 0), ("MalgunGothic", "/mnt/c/Windows/Fonts/malgun.ttf", 0),
    ("NanumGothic", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf", 0),
    ("NotoSansCJK", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 1),
    ("NotoSansCJK", "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc", 1),
    ("AppleGothic", "/System/Library/Fonts/Supplemental/AppleGothic.ttf", 0),
]
_FONT_BOLD = {"MalgunGothic": r"C:\Windows\Fonts\malgunbd.ttf", "NanumGothic": "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"}
_font_cache: dict = {}


def _pdf_fonts() -> tuple[str, str]:
    """(본문 폰트명, 굵은 폰트명) — 한 번 등록한 뒤 재사용.
    우선순위: .env 의 PMC_PDF_FONT(ttf 경로) → config/fonts/*.ttf → 윈도우 맑은 고딕 → 나눔고딕 → 내장 한글 CID 폰트"""
    if _font_cache:
        return _font_cache["r"], _font_cache["b"]
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    reg = bold = None
    cands = list(_FONT_CANDIDATES)
    custom = [Path(p) for p in [os.getenv("PMC_PDF_FONT", "")] if p] + sorted((Path(__file__).resolve().parent.parent / "config" / "fonts").glob("*.ttf"))
    cands = [(f"Custom{i}", str(p), 0) for i, p in enumerate(custom)] + cands
    for name, path, idx in cands:
        if Path(path).exists():
            try:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=idx))
                reg = bold = name
                bp = _FONT_BOLD.get(name)
                if bp and Path(bp).exists():
                    pdfmetrics.registerFont(TTFont(name + "-Bold", bp))
                    bold = name + "-Bold"
                break
            except Exception:
                continue
    if not reg:                                           # 폰트 파일이 없어도 동작(뷰어 내장 한글 폰트 사용)
        pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
        reg = bold = "HYGothic-Medium"
    _font_cache.update(r=reg, b=bold)
    return reg, bold


def _para(text: str) -> str:
    return html.escape(text).replace("\n", "<br/>")


def build_pdf(m: dict, actions: list[dict] | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold = _pdf_fonts()
    S = {
        "proj": ParagraphStyle("proj", fontName=reg, fontSize=8.5, textColor=colors.HexColor("#5B6675"), leading=12),
        "title": ParagraphStyle("title", fontName=bold, fontSize=20, textColor=colors.HexColor("#1F4E9C"), alignment=1, leading=28, spaceAfter=8),
        "k": ParagraphStyle("k", fontName=bold, fontSize=9.5, leading=14),
        "v": ParagraphStyle("v", fontName=reg, fontSize=10, leading=15),
        "h": ParagraphStyle("h", fontName=bold, fontSize=12.5, textColor=colors.HexColor("#1F4E9C"), leading=18, spaceBefore=12, spaceAfter=4),
        "body": ParagraphStyle("body", fontName=reg, fontSize=10.5, leading=17),
        "foot": ParagraphStyle("foot", fontName=reg, fontSize=8, textColor=colors.HexColor("#8A97A8"), leading=11, spaceBefore=16),
    }
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=f"회의록 - {_s(m.get('title'))}", author="KOICA AI 성과관리")
    story = [Paragraph(_para(PROJECT), S["proj"]), Paragraph("회 의 록", S["title"])]
    meta = _meta(m)
    data = [[Paragraph(_para(meta[0][0]), S["k"]), Paragraph(_para(meta[0][1]), S["v"]), "", ""]]
    rest = meta[1:]
    for i in range(0, len(rest), 2):
        pair = rest[i:i + 2]
        row = []
        for k, v in pair:
            row += [Paragraph(_para(k), S["k"]), Paragraph(_para(v), S["v"])]
        row += [""] * (4 - len(row))
        data.append(row)
    tbl = Table(data, colWidths=[30 * mm, 55 * mm, 30 * mm, 55 * mm])
    style = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C2D3")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F9")), ("BACKGROUND", (2, 1), (2, -1), colors.HexColor("#EEF2F9")),
             ("SPAN", (1, 0), (3, 0)), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
    tbl.setStyle(TableStyle(style))
    story += [tbl]
    for title, col in SECTIONS:
        story += [Paragraph(_para(title), S["h"]), Paragraph(_para(_s(m.get(col)) or "(없음)"), S["body"])]
    lines = _action_lines(actions)
    if actions:
        story += [Spacer(1, 6), Paragraph("등록된 후속조치 (현황)", S["k"]), Spacer(1, 3)]
        adata = [[Paragraph(x, S["k"]) for x in ("조치 내용", "담당", "기한", "상태")]]
        for a in actions:
            adata.append([Paragraph(_para(_s(a.get(k))), S["v"]) for k in ("description", "owner", "due_date", "status")])
        at = Table(adata, colWidths=[95 * mm, 30 * mm, 25 * mm, 20 * mm])
        at.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C2D3")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F9")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(at)
    story.append(Paragraph(_para(f"회의ID {_s(m.get('meeting_id')) or '(저장 전)'} · 출력 {datetime.now():%Y-%m-%d %H:%M} · KOICA AI 성과관리"), S["foot"]))
    doc.build(story)
    return buf.getvalue()


# ───────────────────────── 주간업무보고 (주차 × 구분) ─────────────────────────
WEEKLY_HEAD = ["구분", "작성자", "직책", "활동내용"]


def build_weekly_docx(title: str, rows: list[dict], period: str = "") -> bytes:
    """rows: [{role, author, position, content}] — settings 의 열 순서. 빈 칸도 행으로 남긴다."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "맑은 고딕", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    def _font(run, size=None, bold=None, color=None):
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def _shade(cell, hex_color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    p = doc.add_paragraph()
    _font(p.add_run(PROJECT), 9, color="5B6675")
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(h.add_run("주 간 업 무 보 고"), 20, True, "1F4E9C")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(sub.add_run(title + (f"  ·  {period}" if period else "")), 12, True)
    sub.paragraph_format.space_after = Pt(10)

    t = doc.add_table(rows=1, cols=4)
    t.style, t.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
    for c, k in zip(t.rows[0].cells, WEEKLY_HEAD):
        _shade(c, "EEF2F9")
        c.paragraphs[0].paragraph_format.space_before = c.paragraphs[0].paragraph_format.space_after = Pt(3)
        _font(c.paragraphs[0].add_run(k), 10, True)
    for r in rows:
        cells = t.add_row().cells
        _shade(cells[0], "F7F9FC")
        for c, k in zip(cells, ("role", "author", "position")):
            c.paragraphs[0].paragraph_format.space_before = c.paragraphs[0].paragraph_format.space_after = Pt(3)
            _font(c.paragraphs[0].add_run(_s(r.get(k))), 10, k == "role")
        lines = (_s(r.get("content")) or "-").splitlines() or ["-"]
        first = cells[3].paragraphs[0]
        first.paragraph_format.space_before = first.paragraph_format.space_after = Pt(3)
        _font(first.add_run(lines[0]), 10)
        for ln in lines[1:]:
            para = cells[3].add_paragraph()
            para.paragraph_format.space_after = Pt(2)
            _font(para.add_run(ln if ln.strip() else " "), 10)
    for row in t.rows:
        for c, w in zip(row.cells, (2.6, 2.6, 2.4, 9.4)):
            c.width = Cm(w)
    f = doc.add_paragraph()
    f.paragraph_format.space_before = Pt(18)
    _font(f.add_run(f"출력 {datetime.now():%Y-%m-%d %H:%M} · KOICA AI 성과관리"), 8.5, color="8A97A8")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_weekly_pdf(title: str, rows: list[dict], period: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    reg, bold = _pdf_fonts()
    S = {"proj": ParagraphStyle("proj", fontName=reg, fontSize=8.5, textColor=colors.HexColor("#5B6675"), leading=12),
         "title": ParagraphStyle("title", fontName=bold, fontSize=20, textColor=colors.HexColor("#1F4E9C"), alignment=1, leading=28),
         "sub": ParagraphStyle("sub", fontName=bold, fontSize=11.5, alignment=1, leading=16, spaceAfter=10),
         "k": ParagraphStyle("k", fontName=bold, fontSize=9.5, leading=14),
         "v": ParagraphStyle("v", fontName=reg, fontSize=9.5, leading=14.5),
         "foot": ParagraphStyle("foot", fontName=reg, fontSize=8, textColor=colors.HexColor("#8A97A8"), leading=11, spaceBefore=16)}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=title, author="KOICA AI 성과관리")
    story = [Paragraph(_para(PROJECT), S["proj"]), Paragraph("주 간 업 무 보 고", S["title"]),
             Paragraph(_para(title + (f"  ·  {period}" if period else "")), S["sub"])]
    data = [[Paragraph(k, S["k"]) for k in WEEKLY_HEAD]]
    for r in rows:
        data.append([Paragraph(_para(_s(r.get("role"))), S["k"]), Paragraph(_para(_s(r.get("author"))), S["v"]),
                     Paragraph(_para(_s(r.get("position"))), S["v"]), Paragraph(_para(_s(r.get("content")) or "-"), S["v"])])
    tbl = Table(data, colWidths=[26 * mm, 26 * mm, 24 * mm, 98 * mm], repeatRows=1)
    tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C2D3")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF2F9")), ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#F7F9FC")),
                             ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [tbl, Paragraph(_para(f"출력 {datetime.now():%Y-%m-%d %H:%M} · KOICA AI 성과관리"), S["foot"])]
    doc.build(story)
    return buf.getvalue()
